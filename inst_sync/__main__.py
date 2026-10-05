"""Reporte de instalaciones de los add-ins de VENTAS: quien ya instalo, quien esta al dia y quien falta.
Lee los registros que deja cada PC en 0337_SQDCM\\03370_ADDINS_VENTAS\\instalaciones (los escribe el
instalador y el actualizador automatico) y los cruza con "Equipos e integrantes - VENTAS.xlsx".
Uso: python -m inst_sync run   (tambien corre con Actualizar-Todos.cmd)"""
import html
import json
import os
import sys
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
E = html.escape


def compacta(v):
    v = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode().upper()
    return "".join(ch for ch in v if ch.isalnum())


def leer_equipos(ruta, hoja):
    from openpyxl import load_workbook
    ws = load_workbook(ruta, data_only=True, read_only=True)[hoja]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c or "").strip() for c in filas[0]]
    ie, ii = enc.index("Equipo"), enc.index("Integrante")
    out, vistos = [], set()
    for r in filas[1:]:
        eq = str(r[ie] or "").replace("\xa0", " ").strip()
        it = str(r[ii] or "").replace("\xa0", " ").strip()
        if eq and it and compacta(it) not in vistos:
            vistos.add(compacta(it))
            out.append((eq, it))
    return out


def persona(reg, equipos, alias):
    """Integrante del Excel que corresponde a un registro (por alias, nombre de Windows o usuario)."""
    clave = f"{reg.get('pc', '')}__{reg.get('usuario', '')}"
    for k in (clave, str(reg.get("pc", ""))):  # alias por "PC__usuario" o solo por nombre de PC
        for a, v in alias.items():
            if a.upper() == k.upper():
                return v or None  # "" = no asignar a nadie (sale en PCS SIN IDENTIFICAR)
    # Nombre de la PC (en VENTAS cada PC lleva el nombre de su usuario: MARIOSANCHEZ, LAPJOSEROJAS...)
    pc = compacta(reg.get("pc"))
    pc = pc[3:] if pc.startswith("LAP") else pc
    if len(pc) >= 5:
        hits = []
        for _, it in equipos:
            tk = [compacta(t) for t in str(it).split() if compacta(t)]
            if len(tk) < 2:
                continue
            ops = ["".join(tk), "".join(tk[:3])] + [tk[i] + tk[j] for i in range(len(tk)) for j in range(i + 1, len(tk))]
            ops += [tk[0][0] + tk[1][0] + t for t in tk[2:]]  # iniciales: JPGUTIERREZ
            if any(o == pc or (len(pc) >= 12 and o.startswith(pc)) for o in ops):
                hits.append(it)
        if len(hits) == 1:
            return hits[0]
    for cand in (reg.get("nombre"), reg.get("usuario")):
        c = compacta(cand)
        if len(c) < 4:
            continue
        for _, it in equipos:
            ci = compacta(it)
            if c == ci or (len(c) >= 8 and (ci.startswith(c) or c.startswith(ci))):
                return it
    return None  # sin coincidencia clara: sale en PCS SIN IDENTIFICAR con su PC y usuario


def main():
    if len(sys.argv) < 2 or sys.argv[1] != "run":
        print(__doc__)
        return
    cfg = json.loads((RAIZ / "config.json").read_text(encoding="utf-8"))
    addins = Path(os.environ["APPDATA"]) / "Autodesk/Revit/Addins/2025"
    dist = cfg.get("carpeta_distribucion") or str(
        Path((addins / "carpeta-compartida-config.txt").read_text(encoding="utf-8-sig").strip()) / "03370_ADDINS_VENTAS")
    dist = Path(dist)
    vf = dist / "paquete" / "version.txt"
    publicada = vf.read_text(encoding="utf-8").strip() if vf.exists() else "-"
    equipos = leer_equipos(cfg["equipos_xlsx"], cfg.get("equipos_hoja", "Integrantes"))
    alias = cfg.get("alias_pc", {})
    regs = []
    for f in sorted((dist / "instalaciones").glob("*.json")):
        try:
            regs.append(json.loads(f.read_text(encoding="utf-8-sig")))
        except Exception:
            pass
    por_persona, sin_id = {}, []
    for r in regs:
        p = persona(r, equipos, alias)
        if p:
            por_persona.setdefault(p, []).append(r)
        else:
            sin_id.append(r)
    hace7 = (datetime.now() - timedelta(days=7)).strftime("%Y-%m-%d")

    def estado(rs):
        if not rs:
            return "fail", "Sin instalar"
        r = max(rs, key=lambda x: x.get("version", ""))
        if r.get("version") != publicada:
            return "warn", f"Desactualizado ({r.get('version')})"
        if (r.get("ultima_revision") or "") < hace7:
            return "warn", f"Al día, sin revisar desde {r.get('ultima_revision', '')[:10]}"
        return "pass", "Al día"

    filas = [(eq, it, *estado(por_persona.get(it, [])), por_persona.get(it, [])) for eq, it in equipos]
    n = len(filas)
    al_dia = sum(1 for f in filas if f[2] == "pass")
    desact = sum(1 for f in filas if f[2] == "warn")
    sin = sum(1 for f in filas if f[2] == "fail")
    pct = round(100 * al_dia / n) if n else 0
    color = "green" if pct >= 80 else "orange" if pct >= 50 else "red"
    css = (Path(__file__).parent / "estilo.css").read_text(encoding="utf-8")
    ico = {"pass": '<path d="m5 12.5 4.2 4.2L19 7"/>', "warn": '<path d="M12 4 21 20H3L12 4Z"/><path d="M12 10v4"/><path d="M12 17h.01"/>',
           "fail": '<path d="M12 4 21 20H3L12 4Z"/><path d="M12 10v4"/><path d="M12 17h.01"/>',
           "info": '<circle cx="12" cy="12" r="9"/><path d="M12 11v5"/><path d="M12 8h.01"/>'}
    svg = lambda k: f'<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.7" stroke-linecap="round" stroke-linejoin="round">{ico[k]}</svg>'
    lbl = {"pass": "Correcto", "warn": "Atención", "fail": "Requiere atención", "info": "Información"}
    pill = {"pass": "st closed", "warn": "st open", "fail": "st none"}

    def tabla(cab, rows):
        return ('<div class="table-scroll"><table><thead><tr>' + "".join(f"<th>{E(c)}</th>" for c in cab) + "</tr></thead><tbody>"
                + "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows) + "</tbody></table></div>")

    def det(r):
        return (E(persona(r, equipos, alias) or "—"), E(r.get("pc", "")), E(r.get("usuario", "")), E(r.get("version", "")), E(r.get("instalado", "")),
                E(r.get("actualizado", "")), E(r.get("ultima_revision", "")))

    cards = []
    for eq in dict.fromkeys(e for e, _ in equipos):
        fs = [f for f in filas if f[0] == eq]
        k = "pass" if all(f[2] == "pass" for f in fs) else "fail" if all(f[2] == "fail" for f in fs) else "warn"
        ok = sum(1 for f in fs if f[2] == "pass")
        rows = []
        for _, it, st, txt, rs in fs:
            r = max(rs, key=lambda x: x.get("version", "")) if rs else {}
            rows.append((E(it), f'<span class="{pill[st]}">{E(txt)}</span>', E(r.get("pc", "—")), E(r.get("version", "—")),
                         E(r.get("ultima_revision", "—"))))
        cards.append(f'<article class="check-card {k} has-table"><div class="card-top"><div class="status-icon">{svg(k)}</div><div class="card-copy">'
                     f'<div class="status-row"><span class="status-pill">{lbl[k]}</span><span class="result-count">{ok}/{len(fs)} al día</span></div>'
                     f'<h3 class="check-title">{E(eq)}</h3><p class="check-summary">{len(fs)} integrantes · {sum(1 for f in fs if f[2] == "fail")} sin instalar</p></div></div>'
                     + tabla(["Integrante", "Estado", "PC", "Versión", "Última revisión"], rows) + "</article>")
    sin_html = ""
    if sin_id:
        sin_html = (f'<section class="section"><header class="section-heading"><div><p class="section-kicker">Instalaciones · Área de revisión</p><h2 class="section-title">PCS SIN IDENTIFICAR</h2></div><span class="section-count">{len(sin_id)}</span></header>'
                    f'<div class="check-grid"><article class="check-card warn has-table"><div class="card-top"><div class="status-icon">{svg("warn")}</div><div class="card-copy">'
                    '<div class="status-row"><span class="status-pill">Atención</span></div><h3 class="check-title">Instalaciones que no se pudieron relacionar con el Excel</h3>'
                    '<p class="check-summary">PCs cuyo nombre no se pudo relacionar con el Excel. Para asignarla, en INSTALADOR_VENTAS\\config.json agrega dentro de alias_pc: "NOMBREPC": "NOMBRE DEL INTEGRANTE" (o "" para no asignarla a nadie).</p></div></div>'
                    + tabla(["Integrante", "PC", "Usuario Windows", "Versión", "Instalado", "Actualizado", "Última revisión"], [det(r) for r in sin_id]) + "</article></div></section>")
    todos = tabla(["Integrante", "PC", "Usuario Windows", "Versión", "Instalado", "Actualizado", "Última revisión"], [det(r) for r in regs]) if regs else '<p class="empty-mini">Todavía no hay instalaciones registradas.</p>'
    ahora = datetime.now().strftime("%d/%m/%Y %H:%M")
    page = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Instalaciones de add-ins · VENTAS</title><style>{css}
.reveal{{opacity:1;transform:none}} .st{{display:inline-block;padding:3px 9px;border-radius:999px;font-size:11px;font-weight:800}}
.st.closed{{background:rgba(52,199,89,.14);color:#248a3d}} .st.open{{background:rgba(255,149,0,.14);color:#b86e00}} .st.none{{background:rgba(255,59,48,.1);color:#d70015}}
.check-card.warn{{--accent:var(--orange)}} .metric.warn .metric-label{{color:var(--orange)}} .metric.warn .metric-label::before{{background:var(--orange)}}
.empty-mini{{padding:18px;color:var(--tertiary);text-align:center}}</style></head><body><main class="page" style="padding-top:48px">
<section class="hero"><div class="hero-content"><p class="eyebrow"><span class="eyebrow-dot"></span>Add-ins de Revit · VENTAS</p>
<h1><span>Instalaciones.</span><span class="gradient-title health-title">Quién ya tiene los add-ins.</span></h1>
<p class="hero-copy">Finalidad: que todos los integrantes del Excel de equipos tengan los add-ins de VENTAS instalados y en la versión publicada. El porcentaje mide los integrantes al día sobre el total. Meta: 100%.</p>
<div class="file-meta"><div class="file-meta-row"><span class="file-path">Versión publicada: <b>{E(publicada)}</b> · {E(str(dist))}</span></div><div class="file-meta-row"><span>Generado: {ahora}</span></div></div></div>
<div class="score-panel"><div class="score-ring score-{color}" style="--score:{pct}"><div class="score-value"><span class="score-number">{pct}%</span><span class="score-label">AL DÍA</span></div></div>
<p class="score-caption">{al_dia} de {n} integrantes<br>con la versión publicada.</p></div></section>
<section class="metrics"><article class="metric total"><div class="metric-label">Integrantes</div><div class="metric-number">{n}</div></article>
<article class="metric pass"><div class="metric-label">Al día</div><div class="metric-number">{al_dia}</div></article>
<article class="metric warn"><div class="metric-label">Desactualizados / sin revisar</div><div class="metric-number">{desact}</div></article>
<article class="metric fail"><div class="metric-label">Sin instalar</div><div class="metric-number">{sin}</div></article></section>
<section class="section"><header class="section-heading"><div><p class="section-kicker">Instalaciones · Área de revisión</p><h2 class="section-title">EQUIPOS E INTEGRANTES</h2></div><span class="section-count">{len(cards)} checks</span></header><div class="check-grid">{''.join(cards)}</div></section>
{sin_html}
<section class="section"><header class="section-heading"><div><p class="section-kicker">Instalaciones · Área de revisión</p><h2 class="section-title">TODAS LAS INSTALACIONES</h2></div><span class="section-count">{len(regs)} PCs</span></header>
<div class="check-grid"><article class="check-card info has-table"><div class="card-top"><div class="status-icon">{svg("info")}</div><div class="card-copy"><div class="status-row"><span class="status-pill">Detalle</span></div><h3 class="check-title">Registro por PC</h3></div></div>{todos}</article></div></section>
<footer class="footer"><span>Instalaciones · GCPEASA · INSTALADOR_VENTAS</span><span>{ahora}</span></footer></main></body></html>"""
    out = RAIZ / "Dashboard" / "Instalaciones-Ventas-Report.html"
    out.parent.mkdir(exist_ok=True)
    out.write_text(page, encoding="utf-8")
    print(f"Instalaciones: {len(regs)} PCs registradas · {al_dia}/{n} integrantes al dia · {out}")


if __name__ == "__main__":
    main()
